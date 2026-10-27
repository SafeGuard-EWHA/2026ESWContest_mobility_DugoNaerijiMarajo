# -*- coding: utf-8 -*-
"""결선 작품소개서(워드 2p) 생성 — 사무국 양식(docx)에 내용을 채움.

사용: python build_intro.py <양식.docx> <출력.docx>
양식 표 구조(14행): 0 제목/[부문명], 1 팀명, 2 작품명, 3 팀원 머리행, 4~8 팀원,
9 작품 사진, 10 작품 개요, 11 개발 환경, 12 작품 설명, 13 특징 및 장점.
"""
import copy
import os
import sys

import docx
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

SRC, OUT = sys.argv[1], sys.argv[2]
IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), "img")
FONT = "나눔명조"
BLACK = RGBColor(0, 0, 0)
RED = RGBColor(0xC0, 0x00, 0x00)  # 팀이 채워야 할 자리

d = docx.Document(SRC)
t = d.tables[0]


def style_run(r, size=11, bold=False, color=BLACK):
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    rpr = r._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, fonts)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(k), FONT)


def clear_cell(cell):
    """첫 문단 서식(pPr)만 남기고 내용을 비운다."""
    first = cell.paragraphs[0]
    for p in cell.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for child in list(first._p):
        if child.tag != qn("w:pPr"):
            first._p.remove(child)
    return first


def fill(cell, paras, size=11, align=None):
    """paras: 문단 목록. 문단 = str | [(text, {bold, color})...]"""
    first = clear_cell(cell)
    if align is not None:
        first.alignment = align
    for i, para in enumerate(paras):
        if i == 0:
            p = first
        else:
            new = copy.deepcopy(first._p)
            for child in list(new):
                if child.tag != qn("w:pPr"):
                    new.remove(child)
            cell._tc.append(new)
            p = cell.paragraphs[-1]
        runs = para if isinstance(para, list) else [(para, {})]
        for text, opt in runs:
            style_run(p.add_run(text), size=size, bold=opt.get("bold", False), color=opt.get("color", BLACK))


def cell(ri, ci):
    """병합을 고려해 ri행의 ci번째 '서로 다른' 셀."""
    seen, out = set(), []
    for c in t.rows[ri].cells:
        if id(c._tc) not in seen:
            seen.add(id(c._tc)); out.append(c)
    return out[ci]


def todo(text):
    return (text, {"color": RED})


# ---------------- 머리 ----------------
head = cell(0, 0)
for p in head.paragraphs:
    for r in p.runs:
        if r.font.color is not None and r.font.color.type is not None and r.font.color.rgb == RGBColor(0, 0, 0xFF):
            r.text = ""
    if "[" in p.text or p.text.strip() == "":
        pass
# [부문명] 자리(파란 글씨)를 지운 문단에 부문명을 검정으로 다시 씀
target = next(p for p in head.paragraphs if p.text.strip() == "")
style_run(target.add_run("[자동차/모빌리티]"), size=16)

fill(cell(1, 1), ["두고내리지말아조"])
fill(cell(2, 1), ["SafeGuard: UWB와 mmWave 융합 Cascading Wake-up 기반 차량 내 잔류 탑승자 선제 보호 시스템"])

# ---------------- 팀원 (3명, 남는 2행 삭제) ----------------
members = [
    ("팀장", [("윤현서", {})], [("이화여자대학교 ", {}), todo("[학과]")], [("블록 A 전원, 구동, ESP32, 통합", {})]),
    ("팀원", [("최서현", {})], [("이화여자대학교 ", {}), todo("[학과]")], [("블록 C UWB 이탈 판정, 통합", {})]),
    ("팀원", [("윤홍은", {})], [("이화여자대학교 ", {}), todo("[학과]")], [("블록 B 레이더, 카메라, Pi 5, 시뮬레이터", {})]),
]
for k, (role, name, org, job) in enumerate(members):
    ri = 4 + k
    fill(cell(ri, 1), [role], size=10)
    fill(cell(ri, 2), [name], size=10)
    fill(cell(ri, 3), [org], size=10)
    fill(cell(ri, 4), [job], size=10)
# 팀원 사진 (세로 병합된 왼쪽 칸)
tp = clear_cell(cell(3, 0))
tp.add_run().add_picture(os.path.join(IMG, "final", "team.jpg"), width=Cm(4.2))

for ri in (8, 7):
    tr = t.rows[ri]._tr
    tr.getparent().remove(tr)

# ---------------- 작품 사진 ----------------
pic = cell(7, 0)  # 행 2개를 지웠으므로 9 → 7
p = clear_cell(pic)
p.add_run().add_picture(os.path.join(IMG, "benz", "car_side_teddy.jpg"), width=Cm(6.6))
p.add_run(" ")
p.add_run().add_picture(os.path.join(IMG, "final", "system_diagram_crop.png"), width=Cm(7.4))
cap = copy.deepcopy(p._p)
for child in list(cap):
    if child.tag != qn("w:pPr"):
        cap.remove(child)
pic._tc.append(cap)
style_run(pic.paragraphs[-1].add_run("(좌) 시연 시제품: 유아용 전동차와 인형, 운전대 앞 60 GHz 레이더   (우) 시스템 구성도"), size=10)

# ---------------- 본문 4칸 ----------------
fill(cell(8, 1), [
    "SafeGuard는 보호자가 차를 떠난 뒤 차 안에 남겨진 아이, 노약자, 반려동물을 찾아내고, "
    "실내 온도가 위험 수준에 이르기 전에 먼저 보호 동작을 시작하는 차량 내 잔류 탑승자 보호 시스템이다.",
    "보호자가 지닌 UWB 노드로 이탈(5 m 초과)을 확인하면 PIR과 카메라 → mmWave 레이더 → 온도 예측 순서로 "
    "필요한 센서만 차례로 깨우는 Cascading Wake-up 구조를 사용해, 주차 중 전력 소모를 줄이면서 "
    "움직이지 않는 탑승자까지 놓치지 않는다.",
], align=WD_ALIGN_PARAGRAPH.JUSTIFY)
fill(cell(9, 1), [
    [("하드웨어: ", {"bold": True}), ("Raspberry Pi 5(판단, 예측, 로그), ESP32(센서 입출력과 구동), TI IWR6843ISK-ODS 60 GHz 레이더, "
                                    "DWM1001-DEV UWB ×2, Camera Module 3, PIR, SHT31 온습도, 12V 팬과 부저(MOSFET), 비상정지 스위치", {})],
    [("OS와 언어: ", {"bold": True}), ("Raspberry Pi OS 64-bit, C++17(상태기계, 열 관측기), Python 3.11(센서 허브, 대시보드), Arduino C++(ESP32)", {})],
    [("라이브러리와 도구: ", {"bold": True}), ("pyserial, picamera2, numpy, pandas / TI mmWave SDK 3.6.2, Radar Toolbox 4.00, "
                                              "UniFlash / Arduino IDE 2 / Streamlit, pytest(TwinLab 시뮬레이터)", {})],
    [("협업: ", {"bold": True}), ("GitHub Organization(SafeGuard-EWHA), SSH 원격 개발", {})],
], align=WD_ALIGN_PARAGRAPH.JUSTIFY)
fill(cell(10, 1), [
    [("① 이탈 확인 ", {"bold": True}), ("차량측과 보호자측 UWB 거리 이동평균이 5 m를 넘어 10초 이상 유지되면 감지를 시작한다.", {})],
    [("② 1차 감지 ", {"bold": True}), ("PIR 또는 카메라(프레임 차분, 3프레임 연속) 중 하나라도 움직임을 잡으면 레이더 검증으로 넘어간다.", {})],
    [("③ 존재 검증 ", {"bold": True}), ("TI 3D People Tracking 출력(TLV)을 Pi에서 직접 해석해 추적 대상이나 재실 신호가 있으면 탑승자로 확정하고, "
                                        "5초 연속 없으면 공석으로 본다. 추적 영역을 차 실내 크기로 제한해 정지 착석 1인을 95.7% 정확히 판정했다.", {})],
    [("④ 예측과 대응 ", {"bold": True}), ("SHT31 온도의 최근 10분 상승 추세와 적응형 RC 모델로 35°C 도달 시간을 계산해, "
                                        "60초 안에 도달할 것으로 보이면 ESP32가 팬과 부저를 먼저 구동하고 보호자에게 알린다.", {})],
    [("⑤ Fail-Safe ", {"bold": True}), ("레이더 통신이 끊기면 PIR AND 카메라로 보수적으로 판단(SUSPECTED)해 보호를 이어간다.", {})],
], align=WD_ALIGN_PARAGRAPH.JUSTIFY)
fill(cell(11, 1), [
    [("단계적 활성화: ", {"bold": True}), ("앞 단계 이벤트가 있을 때만 다음 센서와 연산을 켜 주차 중 소비전력을 줄인다.", {})],
    [("선제 보호: ", {"bold": True}), ("현재 온도가 아니라 ‘35°C 도달까지 남은 시간’으로 판단해 위험 전에 대응한다.", {})],
    [("놓치지 않는 감지: ", {"bold": True}), ("정상 시 OR로 미검출을 줄이고 레이더로 확인하며, 고장 시 AND로 ‘모름’과 ‘없음’을 구분한다.", {})],
    [("실측 기반 검증: ", {"bold": True}), ("레이더 수신 누락 0, UWB 거리 잡음 3.5 cm, 카메라 조명 오검출 0% 등 실측 원본으로 그린 그래프와 영상을 GitHub에 공개했다.", {})],
    [("확장성: ", {"bold": True}), ("기존 차량에 모듈로 붙일 수 있고, Euro NCAP 2026 어린이 잔류 감지(CPD) 평가와 통학차량 하차 확인 의무에 대응한다.", {})],
], align=WD_ALIGN_PARAGRAPH.JUSTIFY)

# ---------------- 표 아래 파란 안내 문구 삭제 (Word는 표 뒤 빈 문단 하나가 필요) ----------------
body = d.element.body
paras = d.paragraphs
for p in paras[:-1]:
    body.remove(p._p)
for child in list(paras[-1]._p):
    if child.tag != qn("w:pPr"):
        paras[-1]._p.remove(child)

d.save(OUT)
print("saved", OUT)
